from pathlib import Path

from pypdf import PdfReader

from src.domain.models import LoadedDocument
from src.ingestion.base import BaseLoader


class FileLoader(BaseLoader):
    SUPPORTED = {".txt", ".md", ".pdf", ".docx"}

    def supports(self, source_type: str) -> bool:
        return source_type in {"file", "text"}

    def load(self, source: str, **kwargs) -> list[LoadedDocument]:
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {source}")

        suffix = path.suffix.lower()

        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8")
            return [
                LoadedDocument(
                    content=text.strip(),
                    source_type="file",
                    source_uri=str(path),
                    metadata={"file_type": suffix.lstrip(".")},
                )
            ]

        if suffix == ".docx":
            try:
                from docx import Document as DocxDocument
            except ImportError as exc:
                raise ImportError(
                    "Install python-docx to load .docx files: pip install python-docx"
                ) from exc
            doc = DocxDocument(str(path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)
            return [
                LoadedDocument(
                    content=text,
                    source_type="file",
                    source_uri=str(path),
                    metadata={"file_type": "docx", "paragraph_count": len(paragraphs)},
                )
            ]

        if suffix == ".pdf":
            reader = PdfReader(str(path))
            documents: list[LoadedDocument] = []
            for page_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    documents.append(
                        LoadedDocument(
                            content=text.strip(),
                            source_type="file",
                            source_uri=str(path),
                            metadata={"file_type": "pdf", "page": page_num},
                        )
                    )
            return documents

        raise ValueError(f"Unsupported file type: {suffix}")